from __future__ import annotations

from io import BytesIO
import json

from fastapi import (
    APIRouter,
    Depends,
    File,
    HTTPException,
    Response,
    UploadFile,
    status,
)

from docx import Document
from pypdf import PdfReader
from sqlalchemy.orm import Session

from app.api.dependencies import get_current_user
from app.db.session import get_db
from app.models.user import User

from app.schemas.resume import (
    ResumeCreate,
    ResumeResponse,
    ResumeUpdate,
)

from app.services.resume_parser import parse_resume_text

from app.services.resume_service import (
    create_resume,
    delete_resume,
    get_resume_by_id,
    get_user_resumes,
    update_resume,
)


# ======================================================
# ROUTER
# ======================================================

router = APIRouter(
    prefix="/resumes",
    tags=["Resumes"],
)


# ======================================================
# CLEAN TEXT
# ======================================================

def clean_extracted_text(text: str) -> str:
    """
    Clean text extracted from PDF/DOCX/OCR.
    """

    if not text:
        return ""

    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    lines = []

    for line in text.split("\n"):
        line = line.strip()

        if line:
            lines.append(line)

    return "\n".join(lines).strip()


# ======================================================
# NORMAL PDF TEXT EXTRACTION
# ======================================================

def extract_pdf_text(
    file_bytes: bytes,
) -> str:
    """
    Try normal PDF text extraction using pypdf.
    """

    pdf_file = BytesIO(file_bytes)

    reader = PdfReader(pdf_file)

    pages: list[str] = []

    for page in reader.pages:
        page_text = page.extract_text() or ""

        page_text = page_text.strip()

        if page_text:
            pages.append(page_text)

    return clean_extracted_text(
        "\n\n".join(pages)
    )


# ======================================================
# OCR PDF EXTRACTION
# ======================================================

def extract_pdf_text_with_ocr(
    file_bytes: bytes,
) -> str:
    """
    Extract text from scanned/image-based PDFs
    using PyMuPDF + Tesseract OCR.
    """

    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError(
            "PyMuPDF is not installed. "
            "Run: pip install pymupdf"
        ) from exc

    try:
        from PIL import Image
        import pytesseract
        pytesseract.pytesseract.tesseract_cmd = (
            r"C:\Program Files\Tesseract-OCR\tesseract.exe"
        )
    except ImportError as exc:
        raise RuntimeError(
            "pytesseract or Pillow is not installed. "
            "Run: pip install pytesseract pillow"
        ) from exc

    try:
        document = fitz.open(
            stream=file_bytes,
            filetype="pdf",
        )

        pages: list[str] = []

        for page_number, page in enumerate(
            document,
            start=1,
        ):
            print(
                f"OCR processing page {page_number}..."
            )

            # Render PDF page at 2x resolution.
            matrix = fitz.Matrix(
                2.0,
                2.0,
            )

            pixmap = page.get_pixmap(
                matrix=matrix,
                alpha=False,
            )

            image = Image.frombytes(
                "RGB",
                (
                    pixmap.width,
                    pixmap.height,
                ),
                pixmap.samples,
            )

            page_text = pytesseract.image_to_string(
                image,
                config="--psm 6",
            )

            page_text = clean_extracted_text(
                page_text
            )

            if page_text:
                pages.append(page_text)

        document.close()

        return clean_extracted_text(
            "\n\n".join(pages)
        )

    except pytesseract.TesseractNotFoundError as exc:
        raise RuntimeError(
            "Tesseract OCR was not found. "
            "Expected location: "
            r"C:\Program Files\Tesseract-OCR\tesseract.exe"
        ) from exc

    except Exception as exc:
        raise RuntimeError(
            f"OCR failed: {exc}"
        ) from exc


# ======================================================
# DOCX TEXT EXTRACTION
# ======================================================

def extract_docx_text(
    file_bytes: bytes,
) -> str:
    """
    Extract text from DOCX paragraphs and tables.
    """

    docx_file = BytesIO(file_bytes)

    document = Document(docx_file)

    paragraphs: list[str] = []

    # Normal paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    # Tables
    for table in document.tables:
        for row in table.rows:
            cells: list[str] = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    cells.append(cell_text)

            if cells:
                paragraphs.append(
                    " | ".join(cells)
                )

    return clean_extracted_text(
        "\n".join(paragraphs)
    )


# ======================================================
# UPLOAD EXISTING RESUME
# ======================================================

@router.post(
    "/upload",
    response_model=ResumeResponse,
    status_code=status.HTTP_201_CREATED,
)
async def upload_existing_resume(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """
    Upload an existing PDF or DOCX resume.

    PDF handling:

    1. Try normal pypdf extraction.
    2. If no text is found, automatically use OCR.
    3. Parse the extracted text.
    4. Store structured resume data.
    """

    # ==================================================
    # VALIDATE FILE NAME
    # ==================================================

    if not file.filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="No file was provided.",
        )

    filename = file.filename.strip()

    if not filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Invalid filename.",
        )

    lower_filename = filename.lower()

    if not lower_filename.endswith(
        (".pdf", ".docx")
    ):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=(
                "Only PDF and DOCX files are supported."
            ),
        )

    try:
        # ==================================================
        # READ FILE
        # ==================================================

        file_bytes = await file.read()

        if not file_bytes:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="The uploaded file is empty.",
            )

        extracted_text = ""

        # ==================================================
        # PDF
        # ==================================================

        if lower_filename.endswith(".pdf"):

            print(
                f"Processing PDF: {filename}"
            )

            # ----------------------------------------------
            # FIRST: NORMAL TEXT EXTRACTION
            # ----------------------------------------------

            try:
                extracted_text = extract_pdf_text(
                    file_bytes
                )
            except Exception as exc:
                print(
                    "Normal PDF extraction failed:",
                    repr(exc),
                )

                extracted_text = ""

            # ----------------------------------------------
            # SECOND: OCR FALLBACK
            # ----------------------------------------------

            if not extracted_text.strip():

                print(
                    "No normal PDF text found."
                )

                print(
                    "Starting OCR fallback..."
                )

                try:
                    extracted_text = (
                        extract_pdf_text_with_ocr(
                            file_bytes
                        )
                    )

                except Exception as exc:
                    print(
                        "OCR extraction failed:",
                        repr(exc),
                    )

                    raise HTTPException(
                        status_code=(
                            status.HTTP_500_INTERNAL_SERVER_ERROR
                        ),
                        detail=(
                            "The PDF appears to be "
                            "image-based, but OCR "
                            "processing failed. "
                            f"Details: {exc}"
                        ),
                    ) from exc

        # ==================================================
        # DOCX
        # ==================================================

        elif lower_filename.endswith(".docx"):

            print(
                f"Processing DOCX: {filename}"
            )

            extracted_text = extract_docx_text(
                file_bytes
            )

        # ==================================================
        # VALIDATE EXTRACTED TEXT
        # ==================================================

        extracted_text = clean_extracted_text(
            extracted_text
        )

        print(
            "Extracted text length:",
            len(extracted_text),
        )

        if not extracted_text:

            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=(
                    "The file was uploaded successfully, "
                    "but no readable text could be extracted. "
                    "The PDF may be corrupted or contain "
                    "unsupported content."
                ),
            )

        # ==================================================
        # PARSE RESUME
        # ==================================================

        print(
            "Parsing extracted resume text..."
        )

        parsed_resume = parse_resume_text(
            extracted_text
        )

        # ==================================================
        # STORE STRUCTURED RESUME DATA
        # ==================================================

        resume_data = ResumeCreate(
            title=filename,
            content=json.dumps(
                parsed_resume,
                ensure_ascii=False,
            ),
            original_filename=filename,
        )

        resume = create_resume(
            db=db,
            user_id=current_user.id,
            resume_data=resume_data,
        )

        print(
            f"Resume created successfully: ID {resume.id}"
        )

        return resume

    # ==================================================
    # EXPECTED HTTP ERRORS
    # ==================================================

    except HTTPException:
        raise

    # ==================================================
    # UNEXPECTED ERRORS
    # ==================================================

    except Exception as exc:

        print(
            "Resume upload error:",
            repr(exc),
        )

        raise HTTPException(
            status_code=(
                status.HTTP_500_INTERNAL_SERVER_ERROR
            ),
            detail=(
                "Unable to process the uploaded resume."
            ),
        ) from exc


# ======================================================
# CREATE RESUME
# ======================================================

@router.post(
    "",
    response_model=ResumeResponse,
    status_code=status.HTTP_201_CREATED,
)
def create_new_resume(
    resume_data: ResumeCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return create_resume(
        db=db,
        user_id=current_user.id,
        resume_data=resume_data,
    )


# ======================================================
# LIST RESUMES
# ======================================================

@router.get(
    "",
    response_model=list[ResumeResponse],
)
def list_resumes(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    return get_user_resumes(
        db=db,
        user_id=current_user.id,
    )


# ======================================================
# GET RESUME
# ======================================================

@router.get(
    "/{resume_id}",
    response_model=ResumeResponse,
)
def get_resume(
    resume_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    resume = get_resume_by_id(
        db=db,
        resume_id=resume_id,
        user_id=current_user.id,
    )

    if resume is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Resume not found",
        )

    return resume


# ======================================================
# UPDATE RESUME
# ======================================================

@router.patch(
    "/{resume_id}",
    response_model=ResumeResponse,
)
def edit_resume(
    resume_id: int,
    resume_data: ResumeUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    resume = get_resume_by_id(
        db=db,
        resume_id=resume_id,
        user_id=current_user.id,
    )

    if resume is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Resume not found",
        )

    return update_resume(
        db=db,
        resume=resume,
        resume_data=resume_data,
    )


# ======================================================
# DELETE RESUME
# ======================================================

@router.delete(
    "/{resume_id}",
    status_code=status.HTTP_204_NO_CONTENT,
)
def remove_resume(
    resume_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    resume = get_resume_by_id(
        db=db,
        resume_id=resume_id,
        user_id=current_user.id,
    )

    if resume is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Resume not found",
        )

    delete_resume(
        db=db,
        resume=resume,
    )

    return Response(
        status_code=status.HTTP_204_NO_CONTENT
    )